"""Shared, read-only data and file exports for completed teacher exams."""

import csv
from decimal import Decimal
from io import BytesIO, StringIO

from django.db.models import Count
from django.utils import timezone

from results.services import (
    get_conducted_exam_leaderboard,
    get_conducted_exam_summary,
)
from .models import ConductedExamAnswer


HEADERS = (
    'Rank', 'Student', 'Username', 'Joined At', 'Submitted At',
    'Status', 'Score', 'Total Marks', 'Accuracy (%)',
    'Time Taken', 'Violations',
)


def build_exam_report(exam):
    """Use the authoritative ranking and finalized participant data."""
    leaderboard = get_conducted_exam_leaderboard(exam)
    summary = get_conducted_exam_summary(leaderboard)
    participants = [item['participant'] for item in leaderboard]
    total_questions = exam.exam_questions.count()
    total_possible_marks = sum(
        (question.marks or Decimal('0')
         for question in exam.exam_questions.all()),
        Decimal('0'),
    )

    correct_counts = {
        row['participant_id']: row['correct']
        for row in (
            ConductedExamAnswer.objects
            .filter(
                participant_id__in=[p.pk for p in participants],
                is_correct=True,
            )
            .values('participant_id')
            .annotate(correct=Count('pk'))
        )
    }

    rows = []
    for item in leaderboard:
        participant = item['participant']
        student = participant.student
        correct = correct_counts.get(participant.pk, 0)
        accuracy = (
            round(Decimal(correct) * 100 / total_questions, 2)
            if total_questions else None
        )
        rows.append({
            'rank': item['rank'],
            'student': student.get_full_name() or student.username,
            'username': student.username,
            'joined_at': participant.joined_at,
            'submitted_at': participant.submitted_at,
            'status': participant.get_status_display(),
            'score': participant.score,
            'total_marks': participant.total_marks,
            'accuracy': accuracy,
            'time_taken': item['time_taken'],
            'violations': participant.violation_count,
        })

    return {
        'exam': exam,
        'leaderboard': leaderboard,
        'participants': participants,
        'summary': summary,
        'total_students': summary['students'],
        'total_questions': total_questions,
        'total_possible_marks': total_possible_marks,
        'correct_counts': correct_counts,
        'rows': rows,
    }


def _display_date(value):
    return (
        timezone.localtime(value).strftime('%d %b %Y, %I:%M %p')
        if value else '-'
    )


def _safe_text(value):
    """Prevent spreadsheet software from interpreting user text as formulas."""
    text = str(value)
    if text.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + text
    return text


def _export_rows(report):
    for row in report['rows']:
        yield [
            row['rank'],
            _safe_text(row['student']),
            _safe_text(row['username']),
            _display_date(row['joined_at']),
            _display_date(row['submitted_at']),
            row['status'],
            row['score'],
            row['total_marks'],
            row['accuracy'],
            row['time_taken'],
            row['violations'],
        ]


def export_csv(report):
    output = StringIO(newline='')
    writer = csv.writer(output)
    writer.writerow(HEADERS)
    writer.writerows(_export_rows(report))
    return output.getvalue().encode('utf-8-sig')


def export_excel(report):
    """Create an actual XLSX workbook with typed numeric result fields."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Student Results'
    sheet.append(HEADERS)

    for row in _export_rows(report):
        values = list(row)
        for index in (6, 7, 8):
            if values[index] is not None:
                values[index] = float(values[index])
        sheet.append(values)

    header_fill = PatternFill('solid', fgColor='3B2418')
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center', wrap_text=True)

    widths = [9, 30, 24, 24, 24, 20, 15, 15, 17, 20, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
        for index in (6, 7, 8):
            row[index].number_format = '0.00'

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
